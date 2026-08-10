"""Generates the sample documents used by the docling lessons.

Not a lesson itself, just a helper. Run it once from the repo root with
the generation-only dependencies pulled in ad hoc (they are not part of
the project's own pyproject.toml, docling itself never needs them):

    uv run --with reportlab --with python-docx --with python-pptx \
        python lessons/docling/sample_data/_generate/build_samples.py

It writes into lessons/docling/sample_data/, next to this folder.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as RLImage,
)
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OUT = Path(__file__).parent.parent
OUT.mkdir(exist_ok=True)


def build_quarterly_report() -> None:
    """A multi-section PDF with headings, a data table, and a picture.

    This is the workhorse fixture: beginner lessons read its structure,
    intermediate lessons pull the table out with TableFormer and the
    picture out with picture extraction.
    """
    path = OUT / "quarterly_report.pdf"
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=LETTER)
    story = []

    story.append(Paragraph("Q3 Regional Sales Report", styles["Title"]))
    story.append(Spacer(1, 12))
    story.append(Paragraph("Executive Summary", styles["Heading1"]))
    story.append(
        Paragraph(
            "Revenue grew across all three regions this quarter, led by the "
            "North region's expansion into two new distribution partners. "
            "The table below breaks results down by region and product line.",
            styles["BodyText"],
        )
    )
    story.append(Spacer(1, 12))

    story.append(Paragraph("Revenue by Region", styles["Heading2"]))
    table_data = [
        ["Region", "Product Line", "Q2 Revenue", "Q3 Revenue", "Growth"],
        ["North", "Hardware", "$412,000", "$498,000", "+20.9%"],
        ["North", "Services", "$188,000", "$225,000", "+19.7%"],
        ["South", "Hardware", "$276,000", "$291,000", "+5.4%"],
        ["South", "Services", "$142,000", "$150,000", "+5.6%"],
        ["West", "Hardware", "$355,000", "$402,000", "+13.2%"],
        ["West", "Services", "$201,000", "$219,000", "+9.0%"],
    ]
    table = Table(table_data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2b3a55")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")]),
            ]
        )
    )
    story.append(table)
    story.append(PageBreak())

    story.append(Paragraph("Regional Growth Chart", styles["Heading2"]))
    chart_path = _build_bar_chart_image()
    story.append(RLImage(str(chart_path), width=4 * inch, height=2.5 * inch))
    story.append(Spacer(1, 16))

    story.append(Paragraph("Outlook", styles["Heading1"]))
    story.append(
        Paragraph(
            "Q4 guidance assumes the North region's new partners ramp to full "
            "volume by week six. Services revenue is expected to keep pace "
            "with hardware growth as onboarding contracts renew.",
            styles["BodyText"],
        )
    )

    doc.build(story)
    chart_path.unlink()
    print(f"wrote {path}")


def _build_bar_chart_image() -> Path:
    """A tiny bar chart, drawn with PIL so no plotting library is required."""
    img = Image.new("RGB", (600, 360), "white")
    draw = ImageDraw.Draw(img)
    regions = ["North", "South", "West"]
    values = [723, 441, 621]
    max_val = max(values)
    bar_width = 120
    gap = 60
    x = 60
    try:
        font = ImageFont.load_default()
    except Exception:
        font = None
    for region, value in zip(regions, values):
        height = int((value / max_val) * 260)
        top = 300 - height
        draw.rectangle([x, top, x + bar_width, 300], fill="#2b3a55")
        draw.text((x, 305), region, fill="black", font=font)
        draw.text((x, top - 15), f"${value}k", fill="black", font=font)
        x += bar_width + gap
    draw.line([40, 300, 560, 300], fill="black", width=2)
    chart_path = OUT / "_chart_tmp.png"
    img.save(chart_path)
    return chart_path


def build_scanned_invoice() -> None:
    """An image-only PDF: no text layer, exactly what OCR lessons need.

    Rendered as a picture of an invoice and embedded as the entire page,
    the same way a phone-camera or flatbed scan produces a PDF with no
    extractable text.
    """
    img = Image.new("RGB", (850, 1100), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.load_default()
    except Exception:
        font = None
    lines = [
        "INVOICE #4471",
        "",
        "Bill To: Meridian Logistics Co.",
        "Date: 2026-07-14",
        "",
        "Item                         Qty     Unit Price     Total",
        "Warehouse pallet racks       12      $340.00        $4,080.00",
        "Freight handling             1       $650.00        $650.00",
        "Installation labor           18 hrs  $45.00         $810.00",
        "",
        "Subtotal:                                            $5,540.00",
        "Tax (7%):                                            $387.80",
        "Total Due:                                           $5,927.80",
        "",
        "Payment terms: Net 30 days from invoice date.",
    ]
    y = 80
    for line in lines:
        draw.text((80, y), line, fill="black", font=font)
        y += 30

    img_path = OUT / "scanned_invoice.pdf"
    img.save(img_path, "PDF", resolution=150.0)
    print(f"wrote {img_path}")


def build_project_plan_docx() -> None:
    """A DOCX with real Word heading styles, so structure lessons have
    something other than a PDF to compare against.
    """
    doc = Document()
    doc.add_heading("Warehouse Automation Rollout", level=0)
    doc.add_paragraph(
        "This plan covers the phased rollout of the new pick-and-pack "
        "automation across the three regional warehouses."
    )

    doc.add_heading("Phase 1: North Warehouse", level=1)
    doc.add_paragraph(
        "Installation begins the first week of Q4, with staff training "
        "running in parallel during the second week."
    )
    doc.add_paragraph("Conveyor belt replacement", style="List Bullet")
    doc.add_paragraph("Barcode scanner upgrade", style="List Bullet")
    doc.add_paragraph("Staff certification (2 days per shift)", style="List Bullet")

    doc.add_heading("Phase 2: South and West Warehouses", level=1)
    doc.add_paragraph(
        "Rollout follows six weeks after North, once lessons learned from "
        "Phase 1 are folded into the installation checklist."
    )

    doc.add_heading("Risks", level=1)
    doc.add_paragraph(
        "The main risk is parts lead time: conveyor motors currently run "
        "an eight week order-to-delivery window."
    )

    path = OUT / "project_plan.docx"
    doc.save(path)
    print(f"wrote {path}")


def build_team_update_pptx() -> None:
    """A short slide deck, the PPTX counterpart to the DOCX fixture."""
    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]

    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Warehouse Automation: Weekly Update"
    slide.placeholders[1].text = "Operations team, week of July 14"

    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "North Warehouse Status"
    body = slide.placeholders[1].text_frame
    body.text = "Conveyor installation on schedule"
    for line in ["Scanner upgrade complete", "Staff training starts Monday"]:
        p = body.add_paragraph()
        p.text = line
        p.level = 1

    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Open Risks"
    body = slide.placeholders[1].text_frame
    body.text = "Conveyor motor lead time (8 weeks)"
    p = body.add_paragraph()
    p.text = "Mitigation: order South and West motors now"
    p.level = 1

    path = OUT / "team_update.pptx"
    prs.save(path)
    print(f"wrote {path}")


def build_research_note() -> None:
    """A small PDF with a code block and a formula-like line, for the
    formula/code enrichment lesson. Real-world PDFs with LaTeX-rendered
    formulas are hard to fabricate convincingly with reportlab, so this
    keeps expectations honest: enrichment quality depends on how visually
    formula-like the source rendering is.
    """
    path = OUT / "research_note.pdf"
    styles = getSampleStyleSheet()
    mono = styles["Code"]
    doc = SimpleDocTemplate(str(path), pagesize=LETTER)
    story = [
        Paragraph("Note: Batch Size and Convergence", styles["Title"]),
        Spacer(1, 12),
        Paragraph(
            "The loss update for mini-batch gradient descent follows the "
            "standard rule below.",
            styles["BodyText"],
        ),
        Spacer(1, 8),
        Paragraph("theta = theta - eta * gradient(J(theta))", mono),
        Spacer(1, 12),
        Paragraph(
            "A reference implementation of the update step:",
            styles["BodyText"],
        ),
        Spacer(1, 6),
        Paragraph(
            "def sgd_step(theta, grad, eta=0.01):<br/>"
            "&nbsp;&nbsp;&nbsp;&nbsp;return theta - eta * grad",
            mono,
        ),
    ]
    doc.build(story)
    print(f"wrote {path}")


def main() -> None:
    build_quarterly_report()
    build_scanned_invoice()
    build_project_plan_docx()
    build_team_update_pptx()
    build_research_note()


if __name__ == "__main__":
    main()
